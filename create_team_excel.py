import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_railblock_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Team Details"
    
    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Pbm ID",
        "Pbm Stmt",
        "Name of the Students",
        "Gender",
        "Department/School",
        "Program",
        "Specialization",
        "USN",
        "Emailid (Personal)",
        "Emailid (Offical)",
        "Contact Number"
    ]

    pbm_id = "SIH26027"
    pbm_stmt = "AI-Powered Automatic Block Planning to Maximize Asset Availability for Train Operations on Indian Railways"

    data = [
        [
            pbm_id,
            pbm_stmt,
            "Dushyant Acharya",
            "M",
            "Newton School Of Technology",
            "B.Tech",
            "CSE(AI&ML)",
            "2102508745",
            "dushyantacharya873@gmail.com",
            "2102508745@svyasa-sas.edu.in",
            "8988367861"
        ],
        [
            pbm_id,
            pbm_stmt,
            "Pranav Singh",
            "M",
            "Newton School Of Technology",
            "B.Tech",
            "CSE(AI&ML)",
            "2102508781",
            "pranav.nstblr@gmail.com",
            "2102508781@svyasa-sas.edu.in",
            "99005 11200"
        ],
        [
            pbm_id,
            pbm_stmt,
            "Shivansh Goel",
            "M",
            "Newton School Of Technology",
            "B.Tech",
            "CSE(AI&ML)",
            "2102508809",
            "shivanshgoel89@gmail.com",
            "2102508809@svyasa-sas.edu.in",
            "86280 97563"
        ],
        [
            pbm_id,
            pbm_stmt,
            "Tejaswini gangadhara",
            "F",
            "S-VYASA DEEMED TO BE UNIVERSITY",
            "B.Tech",
            "CSE(AI&ML)",
            "2102508316",
            "teju88199@gmail.com",
            "2102508316@svyasa-sas.edu.in",
            "9108629516"
        ],
        [
            pbm_id,
            pbm_stmt,
            "Suchitra K N",
            "F",
            "S-VYASA DEEMED TO BE UNIVERSITY",
            "B.Tech",
            "CSE(AI&ML)",
            "2102508310",
            "suchithrakn14@gmail.com",
            "2102508310@svyasa-sas.edu.in",
            "7483121092"
        ]
    ]

    # Write headers
    ws.append(headers)

    # Write data rows
    for row in data:
        ws.append(row)

    # Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    header_border = Border(
        left=Side(style='thin', color='FFFFFF'),
        right=Side(style='thin', color='FFFFFF'),
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79')
    )

    # Format Header Row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    ws.row_dimensions[1].height = 28

    # Format Data Rows
    align_center_cols = {1, 4, 6, 7, 8, 11} # Pbm ID, Gender, Program, Specialization, USN, Contact
    
    for row_idx in range(2, len(data) + 2):
        ws.row_dimensions[row_idx].height = 22
        # Zebra striping
        row_fill = PatternFill(start_color="F2F5F9" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = thin_border
            
            h_align = "center" if col_idx in align_center_cols else "left"
            cell.alignment = Alignment(horizontal=h_align, vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        
        for cell in col:
            val_str = str(cell.value or '')
            # If line wrap / long string like problem statement, don't make column huge
            if cell.column == 2: # Pbm Stmt
                max_len = 45
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                max_len = max(max_len, len(val_str))
        
        if col[0].column != 2:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        else:
            ws.column_dimensions[col_letter].width = 45

    file_name = "RailBlock-AI.xlsx"
    wb.save(file_name)
    print(f"Successfully generated {file_name}")

if __name__ == "__main__":
    create_railblock_excel()
